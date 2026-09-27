from io import BytesIO, StringIO
from pathlib import Path
import csv
import zipfile

from fastapi import Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from ..database import get_db
from ..imports import import_rows
from ..models import User
from ..security import editor, exporter
from ..services import DAYS, safe_cell, sessions, writable_year
from .common import filters, protected_router

router = protected_router()

EXPORT_COLUMNS = ["school", "teacher", "grade", "class", "day", "start_time", "end_time", "notes"]


def export_rows(entries):
    return [
        [
            safe_cell(x["school_name"]),
            safe_cell(x["teacher_name"]),
            safe_cell(x["grade"]),
            safe_cell(x["class_name"]),
            DAYS[x["day"]],
            x["start_time"],
            x["end_time"],
            safe_cell(x["notes"]),
        ]
        for x in entries
    ]


def spreadsheet_response(rows, fmt, filename):
    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Robotics schedules"
        ws.append(EXPORT_COLUMNS)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in "ABCDEFGH":
            ws.column_dimensions[col].width = 25 if col not in "FG" else 14
        out = BytesIO()
        wb.save(out)
        data = out.getvalue()
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        out = StringIO(newline="")
        writer = csv.writer(out)
        writer.writerow(EXPORT_COLUMNS)
        writer.writerows(rows)
        data = ("﻿" + out.getvalue()).encode("utf-8")
        mime = "text/csv; charset=utf-8"
    return StreamingResponse(
        iter([data]), media_type=mime, headers={"Content-Disposition": f'attachment; filename="{filename}.{fmt}"'}
    )


@router.get("/export")
def export_schedule(
    format: str = Query("csv", pattern="^(csv|xlsx)$"),
    params: dict = Depends(filters),
    user: User = Depends(exporter),
    db: Session = Depends(get_db),
):
    return spreadsheet_response(export_rows(sessions(db, **params)), format, "robotics-schedule")


@router.get("/import/template")
def import_template(user: User = Depends(editor)):
    return spreadsheet_response([], "csv", "schedule-import-template")


@router.post("/import")
def import_schedule(
    academic_year_id: int,
    dry_run: bool = True,
    create_missing: bool = True,
    file: UploadFile = File(...),
    user: User = Depends(editor),
    db: Session = Depends(get_db),
):
    writable_year(db, academic_year_id)
    content = file.file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(413, "Files must be smaller than 5 MB.")
    try:
        suffix = Path(file.filename or "").suffix.lower()
        if suffix == ".csv":
            reader = csv.reader(StringIO(content.decode("utf-8-sig")))
            raw = []
            for row in reader:
                raw.append(row)
                if len(raw) > 2001:
                    raise HTTPException(422, "Import at most 2,000 rows at a time.")
        elif suffix == ".xlsx":
            with zipfile.ZipFile(BytesIO(content)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 30 * 1024 * 1024:
                    raise HTTPException(413, "Expanded workbook exceeds 30 MB.")
            wb = load_workbook(BytesIO(content), read_only=True, data_only=False)
            try:
                raw = list(wb.active.iter_rows(values_only=True, max_row=2002, max_col=20))
                if len(raw) > 2001:
                    raise HTTPException(422, "Import at most 2,000 rows at a time.")
            finally:
                wb.close()
        else:
            raise HTTPException(422, "Upload a .csv or .xlsx file.")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, "The file could not be read. Use the provided CSV or Excel template.")
    if not raw:
        raise HTTPException(422, "The file is empty.")
    header_values = list(raw[0])
    while header_values and header_values[-1] is None:
        header_values.pop()
    headers = [str(x or "").strip().lower() for x in header_values]
    if len(set(headers)) != len(headers) or any(x not in headers for x in EXPORT_COLUMNS[:-1]):
        raise HTTPException(
            422, "Required columns: school, teacher, grade, class, day, start_time, end_time. notes is optional."
        )
    return import_rows(db, user, academic_year_id, raw, headers, create_missing=create_missing, dry_run=dry_run)
