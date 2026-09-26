from datetime import datetime, timedelta
from io import BytesIO, StringIO
from pathlib import Path
import csv
import zipfile
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, FastAPI, File, HTTPException, Query, Request, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select, update
from sqlalchemy.exc import IntegrityError, OperationalError, TimeoutError as PoolTimeoutError
from sqlalchemy.orm import Session
from sqlalchemy.orm.exc import StaleDataError

from .config import get_settings
from .database import get_db
from .imports import import_rows
from .models import (
    AcademicYear,
    AuditLog,
    AuthSession,
    LoginThrottle,
    ScheduleEntry,
    School,
    SchoolClass,
    Teacher,
    TeacherSchoolAssignment,
    User,
    utcnow,
)
from .schemas import (
    AssignmentInput,
    ClassInput,
    LoginInput,
    PasswordInput,
    ResetPassword,
    SchoolInput,
    SessionInput,
    SessionUpdate,
    TeacherInput,
    UserCreate,
    UserUpdate,
    YearInput,
    YearUpdate,
)
from .security import (
    authenticated_user,
    aware,
    create_token,
    current_user,
    dummy_hash,
    editor,
    exporter,
    hash_password,
    super_admin,
    throttle_key,
    verify_password,
)
from .services import (
    DAYS,
    audit,
    commit,
    conflicts,
    find,
    minutes,
    safe_cell,
    sessions,
    snapshot,
    validate_session,
    writable_year,
)

settings = get_settings()
app = FastAPI(
    title="IM-Telligence Robotics Schedule Manager",
    version="1.0.0",
    docs_url=None if settings.production else "/api/docs",
    redoc_url=None,
    openapi_url=None if settings.production else "/api/openapi.json",
)
app.add_middleware(
    CORSMiddleware,
    allow_origins=settings.origins,
    allow_credentials=True,
    allow_methods=["GET", "POST", "PUT", "DELETE"],
    allow_headers=["Content-Type", "X-CSRF-Token", "Authorization"],
)


@app.middleware("http")
async def security_headers(request, call_next):
    origin = request.headers.get("origin")
    if request.method not in {"GET", "HEAD", "OPTIONS"} and origin and origin.rstrip("/") not in settings.origins:
        return JSONResponse(status_code=403, content={"detail": "This origin is not allowed."})
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    response.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=()"
    if request.url.path.startswith("/api"):
        response.headers["Cache-Control"] = "no-store"
    if settings.production:
        response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
        response.headers["Content-Security-Policy"] = (
            "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; font-src 'self'; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
        )
    return response


@app.exception_handler(IntegrityError)
async def integrity_error(_, exc):
    return JSONResponse(
        status_code=409,
        content={
            "detail": "This change conflicts with existing data. Check names, duplicates and teacher assignments."
        },
    )


@app.exception_handler(StaleDataError)
async def stale_error(_, exc):
    return JSONResponse(
        status_code=409, content={"detail": "This schedule entry was modified by another user. Refresh before saving."}
    )


@app.exception_handler(OperationalError)
@app.exception_handler(PoolTimeoutError)
async def database_unavailable(_, exc):
    return JSONResponse(
        status_code=503,
        content={"detail": "The schedule database is temporarily unavailable. Please try again shortly."},
    )


@app.get("/api/health")
def health(db: Session = Depends(get_db)):
    db.execute(select(1))
    return {"status": "ok"}


def user_view(user):
    return {key: value for key, value in snapshot(user).items() if key != "password_hash"}


@app.post("/api/auth/login")
def login(data: LoginInput, response: Response, db: Session = Depends(get_db)):
    email = str(data.email).strip().lower()
    key = throttle_key(email)
    now = utcnow()
    row = db.get(LoginThrottle, key)
    if row is None:
        db.add(LoginThrottle(key=key, attempts=0, window_start=now))
        try:
            db.commit()
        except IntegrityError:
            db.rollback()
    row = db.scalar(select(LoginThrottle).where(LoginThrottle.key == key).with_for_update())
    if aware(row.window_start) < now - timedelta(minutes=15):
        row.attempts = 0
        row.window_start = now
    if row.attempts >= 10:
        raise HTTPException(429, "Too many sign-in attempts. Try again in 15 minutes.")
    row.attempts += 1
    db.commit()
    user = db.scalar(select(User).where(User.email == email))
    valid = verify_password(data.password, user.password_hash if user else dummy_hash)
    if not user or not valid or not user.active:
        raise HTTPException(401, "Email or password is incorrect.")
    db.execute(update(LoginThrottle).where(LoginThrottle.key == key).values(attempts=0))
    user.last_login = now
    token, csrf = create_token(db, user)
    audit(db, user, "sign_in", user)
    commit(db)
    response.set_cookie(
        "schedule_session",
        token,
        httponly=True,
        secure=settings.production,
        samesite="lax",
        max_age=settings.session_hours * 3600,
        path="/api",
    )
    return {"user": user_view(user), "csrf_token": csrf, "access_token": token}


@app.get("/api/auth/me")
def me(request: Request, user: User = Depends(authenticated_user)):
    return {"user": user_view(user), "csrf_token": request.state.claims["csrf"]}


@app.post("/api/auth/logout")
def logout(
    request: Request, response: Response, user: User = Depends(authenticated_user), db: Session = Depends(get_db)
):
    row = db.get(AuthSession, request.state.claims["jti"])
    row.revoked = True
    commit(db)
    response.delete_cookie("schedule_session", path="/api", secure=settings.production, httponly=True, samesite="lax")
    return {"message": "Signed out."}


@app.post("/api/auth/password")
def change_password(
    data: PasswordInput, response: Response, user: User = Depends(authenticated_user), db: Session = Depends(get_db)
):
    if not verify_password(data.current_password, user.password_hash):
        raise HTTPException(422, "Current password is incorrect.")
    if data.current_password == data.new_password:
        raise HTTPException(422, "Choose a different password.")
    old = snapshot(user)
    user.password_hash = hash_password(data.new_password)
    user.must_change_password = False
    user.token_version += 1
    token, csrf = create_token(db, user)
    audit(db, user, "password_changed", user, old)
    commit(db)
    response.set_cookie(
        "schedule_session",
        token,
        httponly=True,
        secure=settings.production,
        samesite="lax",
        max_age=settings.session_hours * 3600,
        path="/api",
    )
    return {"user": user_view(user), "csrf_token": csrf, "access_token": token}


api = APIRouter(prefix="/api", dependencies=[Depends(current_user)])


@api.get("/catalog")
def catalog(db: Session = Depends(get_db)):
    return {
        key: [snapshot(row) for row in db.scalars(select(model).order_by(model.id))]
        for key, model in [
            ("schools", School),
            ("teachers", Teacher),
            ("classes", SchoolClass),
            ("assignments", TeacherSchoolAssignment),
            ("academic_years", AcademicYear),
        ]
    }


@api.get("/schools")
def list_schools(db: Session = Depends(get_db)):
    return [snapshot(row) for row in db.scalars(select(School).order_by(School.name))]


@api.get("/teachers")
def list_teachers(db: Session = Depends(get_db)):
    return [snapshot(row) for row in db.scalars(select(Teacher).order_by(Teacher.full_name))]


def filters(
    academic_year_id: int | None = None,
    school_id: int | None = None,
    teacher_id: int | None = None,
    day: int | None = Query(None, ge=0, le=6),
    grade: str | None = None,
    class_id: int | None = None,
    search: str | None = Query(None, max_length=150),
):
    return dict(
        academic_year_id=academic_year_id,
        school_id=school_id,
        teacher_id=teacher_id,
        day=day,
        grade=grade,
        class_id=class_id,
        search=search,
    )


@api.get("/sessions")
@api.get("/schedules")
def list_sessions(params: dict = Depends(filters), db: Session = Depends(get_db)):
    return sessions(db, **params)


@api.get("/sessions/{session_id}")
@api.get("/schedules/{session_id}")
def get_session(session_id: int, db: Session = Depends(get_db)):
    row = find(db, ScheduleEntry, session_id)
    return next(x for x in sessions(db, academic_year_id=row.academic_year_id) if x["id"] == session_id)


@api.post("/sessions", status_code=201)
@api.post("/schedules", status_code=201)
def create_session(data: SessionInput, user: User = Depends(editor), db: Session = Depends(get_db)):
    validate_session(db, data)
    row = ScheduleEntry(**data.model_dump(), created_by=user.id, updated_by=user.id)
    db.add(row)
    audit(db, user, "create", row)
    commit(db)
    return get_session(row.id, db)


@api.put("/sessions/{session_id}")
@api.put("/schedules/{session_id}")
def update_session(session_id: int, data: SessionUpdate, user: User = Depends(editor), db: Session = Depends(get_db)):
    row = find(db, ScheduleEntry, session_id)
    writable_year(db, row.academic_year_id)
    if row.version != data.expected_version:
        raise HTTPException(409, "This schedule entry was modified by another user. Refresh before saving.")
    validate_session(db, data)
    old = snapshot(row)
    for key, value in data.model_dump(exclude={"expected_version"}).items():
        setattr(row, key, value)
    row.updated_by = user.id
    # Force a versioned UPDATE even when another user submits identical values.
    row.updated_at = utcnow()
    audit(db, user, "update", row, old)
    commit(db)
    return get_session(row.id, db)


@api.delete("/sessions/{session_id}")
@api.delete("/schedules/{session_id}")
def delete_session(
    session_id: int,
    expected_version: int = Query(gt=0),
    user: User = Depends(super_admin),
    db: Session = Depends(get_db),
):
    row = find(db, ScheduleEntry, session_id)
    writable_year(db, row.academic_year_id)
    if row.version != expected_version:
        raise HTTPException(409, "This schedule entry was modified by another user. Refresh before deleting.")
    audit(db, user, "delete", row, snapshot(row))
    db.delete(row)
    commit(db)
    return {"message": "Session deleted."}


@api.get("/conflicts")
def list_conflicts(academic_year_id: int | None = None, db: Session = Depends(get_db)):
    return conflicts(sessions(db, academic_year_id=academic_year_id))


@api.get("/dashboard")
@api.get("/dashboard/stats")
def dashboard(academic_year_id: int, db: Session = Depends(get_db)):
    year = find(db, AcademicYear, academic_year_id)
    entries = sessions(db, academic_year_id=year.id)
    now = datetime.now(ZoneInfo(settings.timezone))
    active_date = year.start_date <= now.date() <= year.end_date
    today = [x for x in entries if active_date and x["day"] == now.weekday()]
    upcoming = []
    for entry in entries:
        distance = (entry["day"] - now.weekday()) % 7
        if distance == 0 and entry["start_time"] <= now.strftime("%H:%M"):
            distance = 7
        next_date = now.date() + timedelta(days=distance)
        while next_date < year.start_date:
            next_date += timedelta(days=7)
        if year.start_date <= next_date <= year.end_date:
            upcoming.append({**entry, "next_date": next_date.isoformat()})
    school_groups, teacher_groups = {}, {}
    for entry in entries:
        for groups, name, color in [
            (school_groups, entry["school_name"], entry["school_color"]),
            (teacher_groups, entry["teacher_name"], entry["school_color"]),
        ]:
            # IDs prevent teachers with identical names from being combined.
            key = entry["school_id"] if groups is school_groups else entry["teacher_id"]
            group = groups.setdefault(key, {"id": key, "name": name, "color": color, "sessions": 0, "minutes": 0})
            group["sessions"] += 1
            group["minutes"] += minutes(entry["end_time"]) - minutes(entry["start_time"])
    return {
        "active_schools": db.scalar(select(func.count()).select_from(School).where(School.active)),
        "teachers": db.scalar(select(func.count()).select_from(Teacher).where(Teacher.active)),
        "weekly_sessions": len(entries),
        "weekly_hours": round(sum(minutes(x["end_time"]) - minutes(x["start_time"]) for x in entries) / 60, 1),
        "conflicts": len(conflicts(entries)),
        "today": today,
        "upcoming": sorted(upcoming, key=lambda x: (x["next_date"], x["start_time"]))[:12],
        "by_school": list(school_groups.values()),
        "by_teacher": list(teacher_groups.values()),
        "date": now.date().isoformat(),
        "timezone": settings.timezone,
    }


@api.post("/schools", status_code=201)
def add_school(data: SchoolInput, user: User = Depends(super_admin), db: Session = Depends(get_db)):
    row = School(**data.model_dump())
    db.add(row)
    audit(db, user, "create", row)
    commit(db)
    return snapshot(row)


@api.put("/schools/{school_id}")
def edit_school(school_id: int, data: SchoolInput, user: User = Depends(super_admin), db: Session = Depends(get_db)):
    row = find(db, School, school_id)
    old = snapshot(row)
    for key, value in data.model_dump().items():
        setattr(row, key, value)
    audit(db, user, "update", row, old)
    commit(db)
    return snapshot(row)


@api.post("/teachers", status_code=201)
def add_teacher(data: TeacherInput, user: User = Depends(super_admin), db: Session = Depends(get_db)):
    row = Teacher(**data.model_dump())
    db.add(row)
    audit(db, user, "create", row)
    commit(db)
    return snapshot(row)


@api.put("/teachers/{teacher_id}")
def edit_teacher(teacher_id: int, data: TeacherInput, user: User = Depends(super_admin), db: Session = Depends(get_db)):
    row = find(db, Teacher, teacher_id)
    old = snapshot(row)
    for key, value in data.model_dump().items():
        setattr(row, key, value)
    audit(db, user, "update", row, old)
    commit(db)
    return snapshot(row)


@api.post("/classes", status_code=201)
def add_class(data: ClassInput, user: User = Depends(editor), db: Session = Depends(get_db)):
    school = find(db, School, data.school_id)
    if not school.active:
        raise HTTPException(422, "Select an active school.")
    row = SchoolClass(**data.model_dump())
    db.add(row)
    audit(db, user, "create", row)
    commit(db)
    return snapshot(row)


@api.post("/assignments", status_code=201)
def assign_teacher(data: AssignmentInput, user: User = Depends(super_admin), db: Session = Depends(get_db)):
    writable_year(db, data.academic_year_id)
    teacher = find(db, Teacher, data.teacher_id)
    school = find(db, School, data.school_id)
    if not teacher.active or not school.active:
        raise HTTPException(422, "Select an active school and teacher.")
    existing = db.scalar(
        select(TeacherSchoolAssignment).where(
            TeacherSchoolAssignment.teacher_id == data.teacher_id,
            TeacherSchoolAssignment.academic_year_id == data.academic_year_id,
        )
    )
    if existing:
        if existing.school_id == data.school_id:
            return snapshot(existing)
        raise HTTPException(
            409,
            "This teacher is already assigned to another school for this academic year. Remove the assignment before reassigning.",
        )
    row = TeacherSchoolAssignment(**data.model_dump())
    db.add(row)
    audit(db, user, "create", row)
    commit(db)
    return snapshot(row)


@api.delete("/assignments/{assignment_id}")
def remove_assignment(assignment_id: int, user: User = Depends(super_admin), db: Session = Depends(get_db)):
    row = find(db, TeacherSchoolAssignment, assignment_id)
    writable_year(db, row.academic_year_id)
    if db.scalar(
        select(ScheduleEntry.id)
        .where(ScheduleEntry.teacher_id == row.teacher_id, ScheduleEntry.academic_year_id == row.academic_year_id)
        .limit(1)
    ):
        raise HTTPException(
            409, "This teacher has sessions in this academic year. Remove those sessions before reassigning."
        )
    audit(db, user, "delete", row, snapshot(row))
    db.delete(row)
    commit(db)
    return {"message": "Assignment removed."}


@api.post("/years", status_code=201)
def add_year(data: YearInput, user: User = Depends(super_admin), db: Session = Depends(get_db)):
    row = AcademicYear(**data.model_dump(), active=False, archived=False)
    db.add(row)
    audit(db, user, "create", row)
    commit(db)
    return snapshot(row)


@api.post("/years/{year_id}/activate")
def activate_year(year_id: int, user: User = Depends(super_admin), db: Session = Depends(get_db)):
    row = writable_year(db, year_id)
    for old_year in db.scalars(select(AcademicYear).where(AcademicYear.active)):
        old = snapshot(old_year)
        old_year.active = False
        old_year.active_slot = None
        audit(db, user, "deactivate", old_year, old)
    db.flush()
    old = snapshot(row)
    row.active = True
    row.active_slot = 1
    audit(db, user, "activate", row, old)
    commit(db)
    return snapshot(row)


@api.put("/years/{year_id}")
def archive_year(year_id: int, data: YearUpdate, user: User = Depends(super_admin), db: Session = Depends(get_db)):
    row = db.scalar(select(AcademicYear).where(AcademicYear.id == year_id).with_for_update())
    if not row:
        raise HTTPException(404, "Academic year not found.")
    if data.archived and row.active:
        raise HTTPException(409, "Activate another academic year before archiving the current year.")
    old = snapshot(row)
    row.archived = data.archived
    audit(db, user, "archive" if data.archived else "restore", row, old)
    commit(db)
    return snapshot(row)


@api.get("/users")
def list_users(user: User = Depends(super_admin), db: Session = Depends(get_db)):
    return [user_view(row) for row in db.scalars(select(User).order_by(User.full_name))]


@api.post("/users", status_code=201)
def add_user(data: UserCreate, user: User = Depends(super_admin), db: Session = Depends(get_db)):
    row = User(
        full_name=data.full_name,
        email=str(data.email).lower(),
        role=data.role,
        export_allowed=data.export_allowed,
        password_hash=hash_password(data.temporary_password),
        must_change_password=True,
    )
    db.add(row)
    audit(db, user, "create", row)
    commit(db)
    return user_view(row)


@api.put("/users/{user_id}")
def edit_user(user_id: int, data: UserUpdate, user: User = Depends(super_admin), db: Session = Depends(get_db)):
    row = find(db, User, user_id)
    if row.id == user.id and (not data.active or data.role != "SUPER_ADMIN"):
        raise HTTPException(409, "You cannot disable or demote your own Super Admin account.")
    old = snapshot(row)
    for key, value in data.model_dump().items():
        setattr(row, key, value)
    if old["role"] != row.role or old["active"] != row.active:
        row.token_version += 1
    audit(db, user, "update", row, old)
    commit(db)
    return user_view(row)


@api.post("/users/{user_id}/reset-password")
def reset_user_password(
    user_id: int, data: ResetPassword, user: User = Depends(super_admin), db: Session = Depends(get_db)
):
    row = find(db, User, user_id)
    if row.id == user.id:
        raise HTTPException(422, "Use Change password for your own account.")
    old = snapshot(row)
    row.password_hash = hash_password(data.temporary_password)
    row.must_change_password = True
    row.token_version += 1
    audit(db, user, "password_reset", row, old)
    commit(db)
    return {"message": "Temporary password set. Existing sessions have been revoked."}


@api.get("/audit")
def list_audit(
    user: User = Depends(super_admin),
    db: Session = Depends(get_db),
    offset: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=200),
):
    q = select(AuditLog, User.full_name).join(User).order_by(AuditLog.timestamp.desc(), AuditLog.id.desc())
    return {
        "items": [{**snapshot(row), "user_name": name} for row, name in db.execute(q.offset(offset).limit(limit))],
        "total": db.scalar(select(func.count()).select_from(AuditLog)),
    }


EXPORT_COLUMNS = ["school", "teacher", "grade", "class", "day", "start_time", "end_time", "notes"]


def export_rows(entries):
    return [
        [
            safe_cell(x["school_name"]),
            safe_cell(x["teacher_name"]),
            safe_cell(x["grade"]),
            safe_cell(x["class_name"]),
            DAYS[x["day"]],
            x["start_time"],
            x["end_time"],
            safe_cell(x["notes"]),
        ]
        for x in entries
    ]


def spreadsheet_response(rows, fmt, filename):
    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Robotics schedules"
        ws.append(EXPORT_COLUMNS)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in "ABCDEFGH":
            ws.column_dimensions[col].width = 25 if col not in "FG" else 14
        out = BytesIO()
        wb.save(out)
        data = out.getvalue()
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        out = StringIO(newline="")
        writer = csv.writer(out)
        writer.writerow(EXPORT_COLUMNS)
        writer.writerows(rows)
        data = ("\ufeff" + out.getvalue()).encode("utf-8")
        mime = "text/csv; charset=utf-8"
    return StreamingResponse(
        iter([data]), media_type=mime, headers={"Content-Disposition": f'attachment; filename="{filename}.{fmt}"'}
    )


@api.get("/export")
def export_schedule(
    format: str = Query("csv", pattern="^(csv|xlsx)$"),
    params: dict = Depends(filters),
    user: User = Depends(exporter),
    db: Session = Depends(get_db),
):
    return spreadsheet_response(export_rows(sessions(db, **params)), format, "robotics-schedule")


@api.get("/import/template")
def import_template(user: User = Depends(editor)):
    return spreadsheet_response([], "csv", "schedule-import-template")


@api.post("/import")
def import_schedule(
    academic_year_id: int,
    dry_run: bool = True,
    create_missing: bool = True,
    file: UploadFile = File(...),
    user: User = Depends(editor),
    db: Session = Depends(get_db),
):
    writable_year(db, academic_year_id)
    content = file.file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(413, "Files must be smaller than 5 MB.")
    try:
        suffix = Path(file.filename or "").suffix.lower()
        if suffix == ".csv":
            reader = csv.reader(StringIO(content.decode("utf-8-sig")))
            raw = []
            for row in reader:
                raw.append(row)
                if len(raw) > 2001:
                    raise HTTPException(422, "Import at most 2,000 rows at a time.")
        elif suffix == ".xlsx":
            with zipfile.ZipFile(BytesIO(content)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 30 * 1024 * 1024:
                    raise HTTPException(413, "Expanded workbook exceeds 30 MB.")
            wb = load_workbook(BytesIO(content), read_only=True, data_only=False)
            try:
                raw = list(wb.active.iter_rows(values_only=True, max_row=2002, max_col=20))
                if len(raw) > 2001:
                    raise HTTPException(422, "Import at most 2,000 rows at a time.")
            finally:
                wb.close()
        else:
            raise HTTPException(422, "Upload a .csv or .xlsx file.")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, "The file could not be read. Use the provided CSV or Excel template.")
    if not raw:
        raise HTTPException(422, "The file is empty.")
    header_values = list(raw[0])
    while header_values and header_values[-1] is None:
        header_values.pop()
    headers = [str(x or "").strip().lower() for x in header_values]
    if len(set(headers)) != len(headers) or any(x not in headers for x in EXPORT_COLUMNS[:-1]):
        raise HTTPException(
            422, "Required columns: school, teacher, grade, class, day, start_time, end_time. notes is optional."
        )
    return import_rows(db, user, academic_year_id, raw, headers, create_missing=create_missing, dry_run=dry_run)


app.include_router(api)

# Optional local bundled preview. Production serves only the API; Vercel serves React.
DIST = Path(__file__).resolve().parents[2] / "frontend" / "dist"
if not settings.production and DIST.exists():

    @app.get("/{path:path}", include_in_schema=False)
    def frontend(path: str):
        if path.startswith("api/"):
            raise HTTPException(404, "API route not found.")
        target = (DIST / path).resolve()
        if target.is_relative_to(DIST.resolve()) and target.is_file():
            return FileResponse(target)
        return FileResponse(DIST / "index.html")
