from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm.exc import StaleDataError

from app.config import Settings
from app.main import app
from app.models import AuditLog, ScheduleEntry, TeacherSchoolAssignment, User
from app.security import verify_password
from conftest import PASSWORD


def create(client, payload):
    r = client.post("/api/schedules", json=payload)
    assert r.status_code == 201, r.text
    return r.json()


def test_authentication_required(store):
    with TestClient(app) as client:
        for path in [
            "/api/schools",
            "/api/teachers",
            "/api/schedules",
            "/api/dashboard/stats?academic_year_id=1",
            "/api/catalog",
            "/api/conflicts",
            "/api/export",
            "/api/audit",
            "/api/users",
        ]:
            assert client.get(path).status_code == 401, path
        assert client.get("/api/health").status_code == 200


def test_multi_user_shared_database(clients, payload):
    entry = create(clients["ADMIN"], payload)
    assert clients["VIEWER"].get("/api/schedules").json()[0]["id"] == entry["id"]
    assert clients["SUPER_ADMIN"].get("/api/schedules").json()[0]["school_color"] == "#3366cc"
    stats = clients["VIEWER"].get("/api/dashboard/stats?academic_year_id=1").json()
    assert stats["weekly_sessions"] == 1
    assert stats["weekly_hours"] == 0.8


def test_viewer_and_admin_permissions(clients, payload):
    viewer = clients["VIEWER"]
    assert viewer.post("/api/schedules", json=payload).status_code == 403
    assert viewer.post("/api/import?academic_year_id=1", files={"file": ("a.csv", b"", "text/csv")}).status_code == 403
    assert viewer.get("/api/export").status_code == 403
    for role in ["VIEWER", "ADMIN"]:
        client = clients[role]
        assert client.get("/api/users").status_code == 403
        assert client.get("/api/audit").status_code == 403
        assert client.post("/api/schools", json={"name": "New", "color": "#112233"}).status_code == 403
        assert (
            client.post(
                "/api/years", json={"name": "Next", "start_date": "2028-09-01", "end_date": "2029-07-31"}
            ).status_code
            == 403
        )
    entry = create(clients["ADMIN"], payload)
    assert clients["ADMIN"].delete(f"/api/schedules/{entry['id']}?expected_version=1").status_code == 403
    assert viewer.put(f"/api/schedules/{entry['id']}", json={**payload, "expected_version": 1}).status_code == 403


def test_teacher_school_rule_and_multiple_teachers(clients, payload):
    admin = clients["SUPER_ADMIN"]
    assert (
        admin.post("/api/assignments", json={"teacher_id": 1, "school_id": 2, "academic_year_id": 1}).status_code == 409
    )
    assert admin.post("/api/schedules", json={**payload, "school_id": 2, "class_id": 3}).status_code == 422
    create(admin, payload)
    create(admin, {**payload, "teacher_id": 2, "class_id": 2})
    assert admin.get("/api/conflicts?academic_year_id=1").json() == []
    assert {x["school_color"] for x in admin.get("/api/schedules").json()} == {"#3366cc"}
    assert admin.delete("/api/assignments/1").status_code == 409
    # The same teacher can belong to a different school in another year.
    create(admin, {**payload, "academic_year_id": 2, "school_id": 2, "class_id": 3})


def test_rule_is_database_enforced(store):
    with store() as db:
        db.add(TeacherSchoolAssignment(teacher_id=1, school_id=2, academic_year_id=1))
        with pytest.raises(IntegrityError):
            db.commit()
        db.rollback()
        db.add(
            ScheduleEntry(
                academic_year_id=1,
                school_id=2,
                teacher_id=1,
                class_id=3,
                day=0,
                start_time="08:00",
                end_time="08:50",
                created_by=1,
                updated_by=1,
            )
        )
        with pytest.raises(IntegrityError):
            db.commit()


def test_optimistic_edit_and_audit(clients, payload, store):
    first = create(clients["ADMIN"], payload)
    url = f"/api/schedules/{first['id']}"
    second = clients["SUPER_ADMIN"].get(url).json()
    edited = clients["ADMIN"].put(url, json={**payload, "start_time": "08:10", "expected_version": first["version"]})
    assert edited.status_code == 200
    assert edited.json()["version"] == 2
    stale = clients["SUPER_ADMIN"].put(
        url, json={**payload, "end_time": "09:00", "expected_version": second["version"]}
    )
    assert stale.status_code == 409
    assert "modified by another user" in stale.json()["detail"]
    assert clients["VIEWER"].get(url).json()["start_time"] == "08:10"
    assert clients["SUPER_ADMIN"].delete(url + "?expected_version=1").status_code == 409
    with store() as db:
        logs = list(db.scalars(select(AuditLog).where(AuditLog.entity_type == "schedule_entries")))
        assert [x.action for x in logs] == ["create", "update"]
        assert logs[-1].old_value["start_time"] == "08:00"
        assert logs[-1].new_value["start_time"] == "08:10"


def test_database_detects_simultaneous_updates(clients, payload, store):
    entry = create(clients["ADMIN"], payload)
    with store() as a, store() as b:
        first = a.get(ScheduleEntry, entry["id"])
        second = b.get(ScheduleEntry, entry["id"])
        first.notes = "First editor"
        a.commit()
        second.notes = "Second editor"
        with pytest.raises(StaleDataError):
            b.commit()


def test_concurrent_api_edits_only_one_succeeds(clients, payload):
    entry = create(clients["ADMIN"], payload)
    url = f"/api/schedules/{entry['id']}"

    def edit(item):
        role, note = item
        return clients[role].put(url, json={**payload, "notes": note, "expected_version": 1}).status_code

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(edit, [("ADMIN", "First"), ("SUPER_ADMIN", "Second")]))
    assert sorted(results) == [200, 409]


def test_conflicts_and_adjacent_sessions(clients, payload):
    admin = clients["ADMIN"]
    create(admin, payload)
    create(admin, {**payload, "class_id": 2, "start_time": "08:30", "end_time": "09:20"})
    create(admin, {**payload, "teacher_id": 2, "start_time": "08:40", "end_time": "09:10"})
    create(admin, {**payload, "class_id": 2, "start_time": "09:20", "end_time": "10:00"})
    items = admin.get("/api/conflicts?academic_year_id=1").json()
    assert len(items) == 2
    assert {kind for x in items for kind in x["types"]} == {"Class overlap", "Teacher overlap"}


@pytest.mark.parametrize(
    "changes", [{"end_time": "07:59"}, {"start_time": "25:00"}, {"day": 7}, {"class_id": 3}, {"unknown": "field"}]
)
def test_server_validation(clients, payload, changes):
    assert clients["ADMIN"].post("/api/schedules", json={**payload, **changes}).status_code == 422


def test_filter_queries_and_export(clients, payload):
    create(clients["ADMIN"], {**payload, "notes": '=HYPERLINK("https://example.com")'})
    create(clients["ADMIN"], {**payload, "teacher_id": 2, "class_id": 2, "day": 1})
    admin = clients["ADMIN"]
    for query in ["teacher_id=1", "day=0", "grade=Grade%204", "class_id=1", "search=Huda"]:
        assert len(admin.get("/api/schedules?" + query).json()) == 1
    assert admin.get("/api/schedules?search=%25").json() == []
    # Repeated values match any of them, e.g. comparing two teachers.
    for query in ["teacher_id=1&teacher_id=2", "day=0&day=1", "grade=Grade%204&grade=Grade%205"]:
        assert len(admin.get("/api/schedules?" + query).json()) == 2
    assert len(admin.get("/api/schedules?teacher_id=1&teacher_id=2&day=1").json()) == 1
    assert admin.get("/api/schedules?day=0&day=7").status_code == 422
    assert admin.get("/api/export?format=csv&teacher_id=1&teacher_id=2").text.count("\n") == 3
    data = admin.get("/api/export?format=csv&teacher_id=1")
    assert data.status_code == 200
    assert "'=HYPERLINK" in data.text
    xlsx = admin.get("/api/export?format=xlsx&teacher_id=1")
    wb = load_workbook(BytesIO(xlsx.content))
    assert wb.active.max_row == 2
    assert wb.active["H2"].data_type == "s"


HEADER = "school,teacher,grade,class,day,start_time,end_time,notes\n"


def test_csv_import_preview_atomicity_and_duplicates(clients):
    admin = clients["ADMIN"]

    def upload(data, dry=True):
        return admin.post(
            f"/api/import?academic_year_id=1&dry_run={str(dry).lower()}",
            files={"file": ("schedule.csv", data, "text/csv")},
        )

    valid = HEADER + "Cedar,Huda,Grade 6,C,Monday,10:00,10:50,Imported\n"
    preview = upload(valid).json()
    assert preview["valid_rows"] == 1 and not preview["errors"]
    assert admin.get("/api/schedules").json() == []
    assert not any(x["grade"] == "Grade 6" for x in admin.get("/api/catalog").json()["classes"])
    invalid = valid + "Horizon,Huda,Grade 7,D,Monday,11:00,11:50,\n"
    failure = upload(invalid, False).json()
    assert len(failure["errors"]) == 1 and failure["imported"] == 0
    assert admin.get("/api/schedules").json() == []
    assert upload(valid, False).json()["imported"] == 1
    duplicate = upload(valid, False).json()
    assert len(duplicate["errors"]) == 1 and duplicate["imported"] == 0


def test_excel_import(clients):
    wb = Workbook()
    wb.active.append(HEADER.strip().split(","))
    wb.active.append(["Cedar", "Rami", "Grade 5", "A", "Tuesday", "10:00", "10:50", ""])
    out = BytesIO()
    wb.save(out)
    r = clients["ADMIN"].post(
        "/api/import?academic_year_id=1&dry_run=false",
        files={"file": ("a.xlsx", out.getvalue(), "application/octet-stream")},
    )
    assert r.status_code == 200, r.text
    assert r.json()["imported"] == 1, r.json()


def test_archives_preserve_data_and_block_writes(clients, payload):
    admin = clients["SUPER_ADMIN"]
    entry = create(admin, payload)
    assert admin.put("/api/years/1", json={"archived": True}).status_code == 409
    assert admin.post("/api/years/2/activate").status_code == 200
    assert admin.put("/api/years/1", json={"archived": True}).status_code == 200
    assert admin.get("/api/schedules?academic_year_id=1").json()[0]["id"] == entry["id"]
    assert admin.put(f"/api/schedules/{entry['id']}", json={**payload, "expected_version": 1}).status_code == 409
    assert admin.post("/api/schedules", json=payload).status_code == 409
    assert admin.get("/api/export?academic_year_id=1").status_code == 200


def test_user_management_passwords_revoke_sessions(clients, store):
    admin = clients["SUPER_ADMIN"]
    with store() as db:
        viewer = db.get(User, 3)
        assert viewer.password_hash.startswith("$argon2")
        assert verify_password(PASSWORD, viewer.password_hash)
    response = admin.post(
        "/api/users",
        json={
            "full_name": "New Viewer",
            "email": "new@example.com",
            "role": "VIEWER",
            "temporary_password": "Temporary-pass-2026",
        },
    )
    assert response.status_code == 201
    assert "password_hash" not in response.json()
    client = TestClient(app)
    auth = client.post("/api/auth/login", json={"email": "new@example.com", "password": "Temporary-pass-2026"}).json()
    client.cookies.clear()
    client.headers["Authorization"] = "Bearer " + auth["access_token"]
    assert client.get("/api/schedules").status_code == 403
    changed = client.post(
        "/api/auth/password", json={"current_password": "Temporary-pass-2026", "new_password": "My-new-password-2026"}
    )
    assert changed.status_code == 200
    client.headers["Authorization"] = "Bearer " + changed.json()["access_token"]
    assert client.get("/api/schedules").status_code == 200
    assert (
        admin.post(
            f"/api/users/{response.json()['id']}/reset-password", json={"temporary_password": "Another-temp-pass-2026"}
        ).status_code
        == 200
    )
    assert client.get("/api/schedules").status_code == 401
    logs = admin.get("/api/audit").json()
    assert "Temporary-pass-2026" not in json.dumps(logs)
    assert "password_hash" not in json.dumps(logs)


def test_disable_user_and_export_permission(clients):
    admin = clients["SUPER_ADMIN"]
    data = {"full_name": "Viewer", "role": "VIEWER", "active": True, "export_allowed": True}
    assert admin.put("/api/users/3", json=data).status_code == 200
    assert clients["VIEWER"].get("/api/export").status_code == 200
    assert admin.put("/api/users/3", json={**data, "active": False}).status_code == 200
    assert clients["VIEWER"].get("/api/schedules").status_code == 401
    assert admin.put("/api/users/1", json={**data, "role": "SUPER_ADMIN", "active": False}).status_code == 409


def test_delete_school_removes_linked_records(clients, store, payload):
    admin = clients["SUPER_ADMIN"]
    create(clients["ADMIN"], payload)
    create(clients["ADMIN"], {**payload, "teacher_id": 2, "class_id": 2})
    assert clients["ADMIN"].delete("/api/schools/1").status_code == 403
    assert clients["ADMIN"].get("/api/schools/1/impact").status_code == 403
    assert admin.get("/api/schools/1/impact").json() == {
        "sessions": 2,
        "assignments": 2,
        "classes": 2,
        "archived_sessions": 0,
    }
    response = admin.delete("/api/schools/1")
    assert response.status_code == 200, response.text
    assert response.json()["removed"] == {"sessions": 2, "assignments": 2, "classes": 2}
    catalog = admin.get("/api/catalog").json()
    assert [x["id"] for x in catalog["schools"]] == [2]
    assert [x["school_id"] for x in catalog["classes"]] == [2]
    assert [x["school_id"] for x in catalog["assignments"]] == [2]
    assert [x["id"] for x in catalog["teachers"]] == [1, 2]  # teachers themselves are kept
    assert admin.get("/api/schedules").json() == []
    assert admin.delete("/api/schools/1").status_code == 404
    with store() as db:
        log = db.scalar(select(AuditLog).where(AuditLog.entity_type == "schools", AuditLog.action == "delete"))
        assert log.old_value["name"] == "Cedar"
        assert len(log.old_value["removed"]["sessions"]) == 2


def test_delete_teacher_removes_sessions_and_assignments_in_every_year(clients, payload):
    admin = clients["SUPER_ADMIN"]
    create(clients["ADMIN"], payload)
    create(clients["ADMIN"], {**payload, "teacher_id": 2, "class_id": 2})
    create(clients["ADMIN"], {**payload, "academic_year_id": 2, "school_id": 2, "class_id": 3})
    assert admin.put("/api/years/2", json={"archived": True}).status_code == 200
    assert admin.get("/api/teachers/1/impact").json() == {"sessions": 2, "assignments": 2, "archived_sessions": 1}
    assert admin.delete("/api/teachers/1").status_code == 200
    remaining = admin.get("/api/schedules").json()
    assert [x["teacher_id"] for x in remaining] == [2]
    catalog = admin.get("/api/catalog").json()
    assert [x["id"] for x in catalog["teachers"]] == [2]
    assert {x["teacher_id"] for x in catalog["assignments"]} == {2}
    assert len(catalog["classes"]) == 3  # classes belong to schools, not teachers


def test_delete_user_keeps_history_and_frees_email(clients, store, payload):
    admin = clients["SUPER_ADMIN"]
    entry = create(clients["ADMIN"], payload)
    assert admin.delete("/api/users/1").status_code == 409  # not yourself
    assert clients["ADMIN"].delete("/api/users/3").status_code == 403
    assert admin.delete("/api/users/2").status_code == 200
    assert clients["ADMIN"].get("/api/schedules").status_code == 401  # existing session revoked
    with TestClient(app) as client:
        login = client.post("/api/auth/login", json={"email": "admin@example.com", "password": PASSWORD})
        assert login.status_code == 401
    assert [x["id"] for x in admin.get("/api/users").json()] == [1, 3]
    assert (
        admin.put(
            "/api/users/2", json={"full_name": "x", "role": "ADMIN", "active": True, "export_allowed": False}
        ).status_code
        == 404
    )
    assert admin.delete("/api/users/2").status_code == 404
    session = admin.get(f"/api/schedules/{entry['id']}").json()
    assert session["updated_by_name"] == "Admin (deleted)"
    names = {x["user_name"] for x in admin.get("/api/audit").json()["items"]}
    assert "Admin (deleted)" in names
    recreated = admin.post(
        "/api/users",
        json={
            "full_name": "Admin Again",
            "email": "admin@example.com",
            "role": "ADMIN",
            "temporary_password": "Temporary-pass-2026",
        },
    )
    assert recreated.status_code == 201
    with store() as db:
        deleted = db.get(User, 2)
        assert deleted.deleted_at is not None and not deleted.active
        assert deleted.email == "deleted-2@deleted.invalid"
        assert not verify_password(PASSWORD, deleted.password_hash)


def test_logout_revokes_bearer(clients):
    admin = clients["ADMIN"]
    assert admin.post("/api/auth/logout").status_code == 200
    assert admin.get("/api/schedules").status_code == 401


def test_cors_and_cookie_csrf(clients, store):
    r = clients["ADMIN"].options(
        "/api/schedules",
        headers={
            "Origin": "https://schedule-test.vercel.app",
            "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "authorization,content-type",
        },
    )
    assert r.status_code == 200
    assert r.headers["access-control-allow-origin"] == "https://schedule-test.vercel.app"
    assert (
        clients["ADMIN"]
        .post(
            "/api/auth/login",
            headers={"Origin": "https://evil.example"},
            json={"email": "admin@example.com", "password": PASSWORD},
        )
        .status_code
        == 403
    )
    with TestClient(app) as cookie_client:
        r = cookie_client.post("/api/auth/login", json={"email": "admin@example.com", "password": PASSWORD})
        assert cookie_client.post("/api/auth/logout").status_code == 403
        assert (
            cookie_client.post("/api/auth/logout", headers={"X-CSRF-Token": r.json()["csrf_token"]}).status_code == 200
        )


def test_login_throttle_and_generic_error(store):
    with TestClient(app) as client:
        for _ in range(10):
            r = client.post("/api/auth/login", json={"email": "missing@example.com", "password": "wrong-password"})
            assert r.status_code == 401
            assert r.json()["detail"] == "Email or password is incorrect."
        assert (
            client.post(
                "/api/auth/login", json={"email": "missing@example.com", "password": "wrong-password"}
            ).status_code
            == 429
        )


def test_production_rejects_sqlite_and_insecure_origins():
    with pytest.raises(ValueError, match="PostgreSQL"):
        Settings(
            app_env="production",
            database_url="sqlite://",
            jwt_secret="b" * 48,
            frontend_url="https://schedule.vercel.app",
        )
    with pytest.raises(ValueError, match="HTTPS"):
        Settings(
            app_env="production",
            database_url="postgresql://example",
            jwt_secret="b" * 48,
            frontend_url="https://schedule.vercel.app,http://insecure.example",
        )
